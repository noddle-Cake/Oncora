"""Schema and quality profile of every raw source (spec 01, Req 2).

Technique: profile before modeling. Column types, null rates, identifier formats and value
ranges decide which joins are possible and which cleaning rules are needed. Writing them down
also turns "I think the data looks like X" into a checked fact.

Access strategy (NFR-3):
- Delimited files go through DuckDB, which streams them, so the 655 MB instinfo.txt is never
  loaded as a DataFrame. Only SQL aggregates come back.
- The CCLE matrix is the exception. DuckDB needs over a minute to parse 19k columns, while
  Polars reads it in a few seconds and it fits easily in memory (32 GB).
"""

import argparse
import re
from dataclasses import dataclass, field
from pathlib import Path

import duckdb
import polars as pl

from oncora.config import DATA_RAW, RAW, REPORTS, raw_path
from oncora.io.checksums import verify_raw

ID_PATTERNS = {
    "DepMap ModelID": r"ACH-\d{6}",
    "Broad BRD id": r"BRD-[A-Z]\d{8}.*",
    "InChIKey": r"[A-Z]{14}-[A-Z]{10}-[A-Z]",
    "Ensembl gene": r"ENSG\d{11}",
    "Cellosaurus": r"CVCL_[A-Z0-9]{4}",
    "ChEMBL id": r"CHEMBL\d+",
    "LINCS SM id": r"LSM-\d+",
}
_ID_MIN_FRACTION = 0.9
_ID_SAMPLE = 2000
_EXAMPLE_WIDTH = 40

# CCLE leading columns that describe the sample rather than a gene.
CCLE_META_COLS = 6


@dataclass
class ColumnProfile:
    name: str
    dtype: str
    null_rate: float
    n_distinct: int
    min: str | None
    max: str | None
    id_pattern: str | None = None
    examples: list[str] = field(default_factory=list)


@dataclass
class TableProfile:
    name: str
    path: str
    n_rows: int
    n_cols: int
    columns: list[ColumnProfile]
    notes: list[str] = field(default_factory=list)


def _clip(v, width: int = _EXAMPLE_WIDTH) -> str | None:
    if v is None:
        return None
    s = str(v)
    return s if len(s) <= width else s[: width - 1] + "…"


def _reader_sql(path: Path, encoding: str = "utf-8") -> str:
    """read_csv call for a raw file. LINCS TSVs write nulls as a literal "" and contain
    backslashes in SMILES, so quoting and escaping are disabled for them."""
    if path.suffix == ".txt":
        return (
            f"read_csv('{path}', delim='\t', header=true, quote='', escape='', "
            f"nullstr=['', '\"\"'], encoding='{encoding}', sample_size=-1)"
        )
    return f"read_csv('{path}', header=true, encoding='{encoding}', sample_size=-1)"


def _detect_id_pattern(con, rel: str, col: str) -> str | None:
    values = [
        r[0]
        for r in con.sql(
            f'SELECT DISTINCT "{col}"::VARCHAR AS v FROM {rel} WHERE "{col}" IS NOT NULL '
            f"ORDER BY hash(v), v LIMIT {_ID_SAMPLE}"
        ).fetchall()
    ]
    if not values:
        return None
    for label, pat in ID_PATTERNS.items():
        rx = re.compile(pat)
        frac = sum(bool(rx.fullmatch(v)) for v in values) / len(values)
        if frac >= _ID_MIN_FRACTION:
            return label if frac == 1 else f"{label} ({frac:.1%})"
    return None


def profile_delimited(
    name: str, path: Path, con=None, encoding: str = "utf-8", notes: list[str] | None = None
) -> TableProfile:
    con = con or duckdb.connect()
    # Materialize once: a view would re-parse the file for every per-column query below.
    con.sql(f"CREATE OR REPLACE TEMP TABLE t AS SELECT * FROM {_reader_sql(path, encoding)}")
    summary = con.sql("SUMMARIZE t").fetchall()
    cols = [r[0] for r in summary]
    # SUMMARIZE's distinct count is approximate and its null % is rounded, so take exact counts.
    count_sql = ", ".join(f'COUNT(DISTINCT "{c}"), COUNT("{c}")' for c in cols)
    n_rows, *counts = con.sql(f"SELECT COUNT(*), {count_sql} FROM t").fetchone()

    profiles = []
    for (col, dtype, mn, mx, *_), n_distinct, n_present in zip(summary, counts[::2], counts[1::2]):
        is_text = dtype == "VARCHAR"
        examples = [
            _clip(r[0])
            for r in con.sql(
                # hash order: deterministic across runs but not just the smallest values
                f'SELECT DISTINCT "{col}" AS v FROM t WHERE v IS NOT NULL '
                "ORDER BY hash(v), v LIMIT 3"
            ).fetchall()
        ]
        profiles.append(
            ColumnProfile(
                name=col,
                dtype=dtype,
                null_rate=1 - n_present / n_rows if n_rows else 0.0,
                n_distinct=n_distinct,
                min=_clip(mn),
                max=_clip(mx),
                id_pattern=_detect_id_pattern(con, "t", col) if is_text else None,
                examples=examples,
            )
        )
    return TableProfile(
        name, str(path.relative_to(DATA_RAW)), n_rows, len(cols), profiles, list(notes or [])
    )


def load_ccle(path: Path | None = None) -> pl.DataFrame:
    return pl.read_csv(path or raw_path("ccle"))


def profile_ccle(df: pl.DataFrame, path: Path) -> TableProfile:
    meta = df.columns[:CCLE_META_COLS]
    genes = df.columns[CCLE_META_COLS:]
    profiles = []
    for col in meta:
        s = df[col]
        profiles.append(
            ColumnProfile(
                name=col or "(unnamed index)",
                dtype=str(s.dtype),
                null_rate=s.null_count() / len(s),
                n_distinct=s.n_unique(),
                min=_clip(s.min()),
                max=_clip(s.max()),
                id_pattern=next(
                    (
                        label
                        for label, pat in ID_PATTERNS.items()
                        if s.dtype == pl.String
                        and s.drop_nulls().str.contains(f"^{pat}$").all()
                    ),
                    None,
                ),
                examples=[_clip(v) for v in s.unique(maintain_order=True).head(3)],
            )
        )

    values = df.select(genes)
    stds = values.std()
    std_row = stds.row(0)
    n_null = sum(values.null_count().row(0))
    gene_name_ok = sum(bool(re.fullmatch(r".+ \(\d+\)", g)) for g in genes)
    notes = [
        f"{len(genes):,} gene columns named `SYMBOL (ENTREZID)` "
        f"({gene_name_ok:,} match the pattern); all {values.dtypes[0]}.",
        f"Gene values: min {values.min().min_horizontal()[0]:.3f}, "
        f"max {values.max().max_horizontal()[0]:.3f}, {n_null:,} nulls (log1p TPM).",
        f"Per-gene std across samples: median {pl.Series(std_row).median():.3f}; "
        f"{sum(s == 0 for s in std_row):,} genes have zero variance.",
        f"{df['ModelID'].n_unique():,} distinct ModelIDs in {df.height:,} rows; "
        f"{(df['IsDefaultEntryForModel'] == 'Yes').sum():,} rows are the default entry "
        "for their model.",
    ]
    return TableProfile(
        "CCLE expression", str(path.relative_to(DATA_RAW)), df.height, df.width, profiles, notes
    )


def profile_xlsx(name: str, path: Path) -> TableProfile:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True)
    notes = [f"Sheet `{ws.title}`: {ws.max_row:,} rows × {ws.max_column} cols" for ws in wb]
    return TableProfile(name, str(path.relative_to(DATA_RAW)), 0, 0, [], notes)


@dataclass
class Check:
    name: str
    expected: str
    measured: str
    ok: bool


def grounding_checks(ccle: pl.DataFrame, con=None) -> list[Check]:
    """Compare measured values with the facts recorded in steering/data-sources.md.

    A failing check is reported, not raised: finding drift is the point of EDA (design §Error Handling).
    """
    con = con or duckdb.connect()
    dss3 = raw_path("dss3")
    n, n_cmp, n_cell, med, p90, mx, zero = con.sql(
        f"""SELECT COUNT(*), COUNT(DISTINCT pert_id), COUNT(DISTINCT cell_line),
                   MEDIAN(sensitivity_score), QUANTILE_CONT(sensitivity_score, 0.9),
                   MAX(sensitivity_score), AVG((sensitivity_score = 0)::INT)
            FROM read_csv('{dss3}', header=true)"""
    ).fetchone()
    dss3_cells = {r[0] for r in con.sql(f"SELECT DISTINCT cell_line FROM '{dss3}'").fetchall()}
    overlap = len(dss3_cells & set(ccle["ModelID"].to_list()))

    def exact(name, expected, measured):
        return Check(name, f"{expected:,}", f"{measured:,}", expected == measured)

    def approx(name, expected, measured, rel=0.05):
        return Check(
            name, f"≈ {expected:g}", f"{measured:.4g}", abs(measured - expected) <= rel * expected
        )

    return [
        exact("CCLE rows", 1775, ccle.height),
        exact("CCLE columns", 19221, ccle.width),
        exact("CCLE distinct ModelID", 1719, ccle["ModelID"].n_unique()),
        exact("CCLE default-entry rows", 1719, int((ccle["IsDefaultEntryForModel"] == "Yes").sum())),
        exact("DSS3 rows", 512950, n),
        exact("DSS3 compounds", 972, n_cmp),
        exact("DSS3 cell lines", 1204, n_cell),
        approx("DSS3 median", 0.087, med),
        approx("DSS3 p90", 0.41, p90),
        approx("DSS3 max", 96, mx),
        approx("DSS3 fraction exactly 0", 0.20, zero, rel=0.1),
        exact("Cell lines in both CCLE and DSS3", 985, overlap),
    ]


def _md_escape(v) -> str:
    return "" if v is None else str(v).replace("|", "\\|").replace("\n", " ")


def render_profile_md(profiles: list[TableProfile], checks: list[Check], checksums_ok: int) -> str:
    out = [
        "# Raw data profile",
        "",
        "Generated by `uv run python -m oncora.io.profile`. Do not edit by hand.",
        "",
        f"Raw-file integrity: **{checksums_ok} files match `data/raw/SHA256SUMS`.**",
        "",
        "## Grounding checks",
        "",
        "Measured values compared with `.kiro/steering/data-sources.md`.",
        "",
        "| Check | Expected | Measured | |",
        "|---|---|---|---|",
    ]
    out += [f"| {c.name} | {c.expected} | {c.measured} | {'✅' if c.ok else '❌'} |" for c in checks]
    out += ["", "## Overview", "", "| Source | Path | Rows | Cols |", "|---|---|---|---|"]
    out += [
        f"| {p.name} | `{p.path}` | {p.n_rows:,} | {p.n_cols:,} |" if p.n_rows else
        f"| {p.name} | `{p.path}` | — | — |"
        for p in profiles
    ]
    for p in profiles:
        out += ["", f"## {p.name}", "", f"`data/raw/{p.path}`", ""]
        out += [f"- {n}" for n in p.notes]
        if p.notes:
            out.append("")
        if p.columns:
            out += [
                "| Column | Type | Null % | Distinct | Min | Max | ID format | Examples |",
                "|---|---|---|---|---|---|---|---|",
            ]
            for c in p.columns:
                out.append(
                    f"| `{_md_escape(c.name)}` | {c.dtype} | {c.null_rate:.1%} | {c.n_distinct:,} "
                    f"| {_md_escape(c.min)} | {_md_escape(c.max)} | {c.id_pattern or ''} "
                    f"| {_md_escape('; '.join(map(str, c.examples)))} |"
                )
    return "\n".join(out) + "\n"


DELIMITED_SOURCES = {
    "DSS3 sensitivity": "dss3",
    "compoundinfo": "compoundinfo",
    "pertinfo": "pertinfo",
    "cellinfo": "cellinfo",
    "geneinfo": "geneinfo",
    "Small-molecule metadata": "sm_metadata",
    "instinfo": "instinfo",
}

# Small_Molecule_Metadata.txt is ASCII except for a stray 0xBF byte in four
# sm_alternative_name cells (tetracaine synonyms), which is invalid UTF-8. Reading it as
# Latin-1 is lossless for every other byte.
SOURCE_ENCODING = {"sm_metadata": "latin-1"}
SOURCE_NOTES = {
    "sm_metadata": [
        "Not valid UTF-8: 4 `sm_alternative_name` cells (tetracaine synonyms, lines "
        "25072 and 26235–26237) contain a stray 0xBF byte. Every other line is ASCII, so the "
        "file is read as Latin-1."
    ]
}


def build_profile() -> str:
    verified = verify_raw()
    con = duckdb.connect()
    ccle = load_ccle()
    profiles = [profile_ccle(ccle, raw_path("ccle"))]
    profiles += [
        profile_delimited(
            n, raw_path(k), con, SOURCE_ENCODING.get(k, "utf-8"), SOURCE_NOTES.get(k)
        )
        for n, k in DELIMITED_SOURCES.items()
    ]
    profiles.append(profile_xlsx("Field definitions", raw_path("field_definitions")))
    return render_profile_md(profiles, grounding_checks(ccle, con), len(verified))


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--out", type=Path, default=REPORTS / "eda_profile.md")
    args = ap.parse_args()
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(build_profile())
    print(f"Wrote {args.out}")


if __name__ == "__main__":
    main()
